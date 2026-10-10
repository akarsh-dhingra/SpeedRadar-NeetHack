"""File tools: inventory, spreadsheet, PDF (text parser or LLM reader)."""

from __future__ import annotations

import re
from pathlib import Path

import neatlogs
from openpyxl import load_workbook
from pydantic import BaseModel

from spendradar import llm

MONEY = re.compile(r"(-?[\d,]+\.\d{2})\s*$")


class ToolError(Exception):
    """A tool could not produce a result; the loop decides how to recover."""


@neatlogs.span(kind="TOOL", name="act:list_files", tool_name="list_files")
def list_files(folder: str) -> list[dict]:
    root = Path(folder)
    return [{"path": str(p.relative_to(root)), "bytes": p.stat().st_size, "suffix": p.suffix.lower()}
            for p in sorted(root.rglob("*")) if p.is_file() and not p.name.startswith(".")]


@neatlogs.span(kind="TOOL", name="act:read_spreadsheet", tool_name="read_spreadsheet")
def read_spreadsheet(path: str, sheet: str, first_row: int = 1, last_row: int | None = None) -> list[dict]:
    """Rows as dicts keyed by the header row, each with its 1-based sheet row number."""
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[sheet]
    except Exception as e:
        raise ToolError(f"cannot open {path}!{sheet}: {e}") from e
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "category")
    header = [str(h) if h is not None else "" for h in rows[header_idx]]
    out = []
    for i, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if i < first_row or (last_row and i > last_row) or not r or r[0] is None:
            continue
        out.append({"_row": i, **{h: v for h, v in zip(header, r)}})
    return out


def sheet_title(path: str, sheet: str) -> list[str]:
    """Non-empty cells of the sheet's first row (the title line above the header)."""
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        first = next(wb[sheet].iter_rows(values_only=True, max_row=1), ())
    except Exception as e:
        raise ToolError(f"cannot open {path}!{sheet}: {e}") from e
    return [str(v) for v in first if v not in (None, "")]


class InvoiceLine(BaseModel):
    description: str
    amount: float


class InvoiceExtract(BaseModel):
    invoice_number: str
    vendor: str
    cost_centre: str
    invoice_date: str
    lines: list[InvoiceLine]
    stated_total: float


def _parse_invoice_text(text: str) -> InvoiceExtract:
    fields: dict[str, str] = {}
    lines: list[InvoiceLine] = []
    total = None
    in_items = False
    all_lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    for ln in all_lines:
        if ":" in ln and not in_items:
            k, _, v = ln.partition(":")
            fields[k.strip().lower()] = v.strip()
        if ln.startswith("Description"):
            in_items = True
            continue
        m = MONEY.search(ln)
        if ln.upper().startswith("TOTAL") and m:
            total = float(m.group(1).replace(",", ""))
            in_items = False
        elif in_items and m:
            lines.append(InvoiceLine(description=ln[: m.start()].strip(), amount=float(m.group(1).replace(",", ""))))
    if total is None or not lines:
        raise ToolError("text parser found no line items / total")
    return InvoiceExtract(invoice_number=fields.get("invoice no.", ""), vendor=all_lines[0],
                          cost_centre=fields.get("cost centre", ""), invoice_date=fields.get("invoice date", ""),
                          lines=lines, stated_total=total)


def _open_error(p: Path) -> str | None:
    """Why the PDF structure is unreadable (no /Root, bad xref, not a PDF), or None if it opens."""
    import pdfplumber

    try:
        with pdfplumber.open(p):
            return None
    except Exception as e:
        return f"{type(e).__name__}: {str(e)[:120]}"


@neatlogs.span(kind="TOOL", name="act:read_pdf", tool_name="read_pdf")
def read_pdf(path: str, mode: str = "text") -> dict:
    """mode='text': pdfplumber + deterministic parse. mode='llm': Gemini reads the PDF bytes."""
    p = Path(path)
    if mode == "text":
        import pdfplumber

        try:
            with pdfplumber.open(p) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        except Exception as e:
            raise ToolError(f"pdfplumber could not open file: {type(e).__name__}: {str(e)[:120]}") from e
        if not text.strip():
            raise ToolError("no text layer (likely a scanned image)")
        inv = _parse_invoice_text(text)
        return {"extract": inv.model_dump(), "parser": "pdfplumber", "llm": None}
    if mode == "llm":
        try:
            data = p.read_bytes()
        except OSError as e:
            raise ToolError(f"cannot read file: {type(e).__name__}: {e}") from e
        # A file the PDF parser cannot open is corrupt, not scanned: the LLM API rejects it with
        # 400 INVALID_ARGUMENT, so fail here without spending a call. Scanned PDFs open fine.
        if err := _open_error(p):
            raise ToolError(f"corrupt PDF, not sent to the LLM reader: {err}")
        try:
            res = llm.complete(
                "extract",
                "Extract this invoice exactly as printed. Amounts as numbers without separators. "
                "Do not compute or correct anything; copy the stated total as printed.",
                schema=InvoiceExtract, attachments=[(data, "application/pdf")])
        except llm.LLMError as e:
            raise ToolError(f"LLM PDF reader failed: {str(e)[:160]}") from e
        return {"extract": res.parsed.model_dump(), "parser": f"llm:{res.model}",
                "llm": {"model": res.model, "tokens_in": res.tokens_in, "tokens_out": res.tokens_out,
                        "latency_s": res.latency_s, "cached": res.cached}}
    raise ToolError(f"unknown mode {mode!r}")
