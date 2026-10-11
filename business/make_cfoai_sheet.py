"""Build the cfo.ai input sheet from business/inputs.md values.

Writes business/cfoai/spend_radar_inputs.xlsx (Assumptions + Monthly sheets) and CSV copies.
Values tagged SAMPLE are placeholders chosen for the plan, not facts; change them in
ASSUMPTIONS below and re-run: uv run python business/make_cfoai_sheet.py
"""

import csv
import math
from pathlib import Path

from openpyxl import Workbook

OUT = Path("business/cfoai")

# (key, value, unit, type, source)
ASSUMPTIONS = [
    ("price_starter", 99, "USD/month", "PROPOSAL", "inputs.md §3"),
    ("price_team", 399, "USD/month", "PROPOSAL", "inputs.md §3"),
    ("runs_starter", 8, "runs/customer/month", "ASSUMPTION", "inputs.md §2"),
    ("runs_team", 40, "runs/customer/month", "ASSUMPTION", "inputs.md §2"),
    ("llm_per_run", 0.0285, "USD/run", "MEASURED", "unit_costs.md, mean of 5 evidence runs"),
    ("crustdata_per_cold_run_base", 9.24, "USD/run", "MEASURED x PUBLISHED", "5 enrich x $1.20 + 360 jobs x $9/1k (high end)"),
    ("crustdata_per_cold_run_growth", 6.24, "USD/run", "MEASURED x PUBLISHED", "5 enrich x $0.60 + 360 jobs x $9/1k (low end)"),
    ("cached_run_share", 0.5, "fraction", "ASSUMPTION", "half the runs reuse cached competitor data"),
    ("hosting", 15.20, "USD/month", "PUBLISHED (third-party)", "t3.small, instances.vantage.sh"),
    ("founder_salary", 2000, "USD/month", "SAMPLE", "placeholder; set your own"),
    ("tools_saas", 100, "USD/month", "SAMPLE", "placeholder: neatlogs, domain, misc"),
    ("starting_cash", 25000, "USD", "SAMPLE", "placeholder; set your own"),
]
A = {k: v for k, v, *_ in ASSUMPTIONS}

# Paying customers at the end of each month (Starter, Team); linear between points.
TARGETS = {
    "Base": {3: (0, 0), 6: (5, 1), 12: (15, 3), 24: (40, 10)},
    "Growth": {3: (0, 0), 6: (10, 2), 12: (35, 8), 24: (100, 25)},
}
CRUST = {"Base": A["crustdata_per_cold_run_base"], "Growth": A["crustdata_per_cold_run_growth"]}


def customers(points: dict, month: int) -> tuple[int, int]:
    if month <= 3:
        return 0, 0
    keys = sorted(points)
    for lo, hi in zip(keys, keys[1:]):
        if lo <= month <= hi:
            f = (month - lo) / (hi - lo)
            # round half up (22.5 -> 23), as cfo.ai does; Python's round() would give 22
            return tuple(math.floor(a + (b - a) * f + 0.5) for a, b in zip(points[lo], points[hi]))
    raise ValueError(month)


def monthly_rows() -> list[dict]:
    rows = []
    for scen, points in TARGETS.items():
        cash = A["starting_cash"]
        cold_cost = CRUST[scen] * (1 - A["cached_run_share"]) + A["llm_per_run"]
        for m in range(1, 25):
            s, t = customers(points, m)
            revenue = s * A["price_starter"] + t * A["price_team"]
            runs = s * A["runs_starter"] + t * A["runs_team"]
            cogs = runs * cold_cost + A["hosting"]
            opex = A["founder_salary"] + A["tools_saas"]
            burn = cogs + opex - revenue
            cash -= burn
            rows.append({
                "scenario": scen, "month": m, "starter_customers": s, "team_customers": t,
                "runs": runs, "revenue": round(revenue, 2), "cogs": round(cogs, 2),
                "gross_profit": round(revenue - cogs, 2),
                "gross_margin": round((revenue - cogs) / revenue, 3) if revenue else "",
                "opex": opex, "net_burn": round(burn, 2), "ending_cash": round(cash, 2),
                "runway_months": round(cash / burn, 1) if burn > 0 else "cash-flow positive",
            })
    return rows


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    rows = monthly_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    ws.append(["key", "value", "unit", "type", "source"])
    for r in ASSUMPTIONS:
        ws.append(list(r))
    ws2 = wb.create_sheet("Monthly")
    ws2.append(list(rows[0]))
    for r in rows:
        ws2.append(list(r.values()))
    wb.save(OUT / "spend_radar_inputs.xlsx")
    with open(OUT / "assumptions.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["key", "value", "unit", "type", "source"])
        w.writerows(ASSUMPTIONS)
    with open(OUT / "monthly.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    for scen in TARGETS:
        for m in (6, 12, 24):
            r = next(x for x in rows if x["scenario"] == scen and x["month"] == m)
            print(scen, m, {k: r[k] for k in ("starter_customers", "team_customers", "revenue", "cogs", "gross_margin", "net_burn", "ending_cash", "runway_months")})


if __name__ == "__main__":
    main()
