"""Generate HELD-OUT fictional companies to test that the agent generalises beyond Acme.

Each company has its own industry, real competitors, headcount mix, vendors and error placements.
Error types beyond the Acme set: a duplicate invoice, a zero-byte PDF, a mis-categorised invoice from a
vendor that appears only once, and a negative control (a 0.4% rounding difference inside the 1%
tolerance that must NOT be flagged). The agent was never tuned on these folders.

Writes data/<key>/ and evals/ground_truth_<key>.json for each company.
Run: uv run python scripts/make_heldout_data.py [key ...]
"""

from __future__ import annotations

import csv
import json
import random
import shutil
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).parent))
from make_acme_data import MONTHS, write_scanned_pdf, write_text_pdf  # noqa: E402  (PDF drawing only)

OBS_PEERS = yaml.safe_load(Path("data/acme/competitors.yaml").read_text(encoding="utf-8"))

# Vendor → (line items as (description, share of base)), monthly base range
VENDORS = {
    "Stratus Hosting Co": ([("Virtual machines", .55), ("Object storage", .25), ("Load balancing + bandwidth", .20)], (55_000, 66_000)),
    "Quillfire Ads": ([("Social ad placements", .7), ("Creative production", .3)], (14_000, 20_000)),
    "Vectorbay Compute": ([("GPU cluster hours", .75), ("Model hosting", .25)], (12_000, 16_000)),
    "Ledgerline CRM": ([("CRM licences", .84), ("Call recording add-on", .16)], (3_700, 3_800)),
    "Corvane Cloud": ([("Compute", .5), ("Managed Postgres", .3), ("Data transfer", .2)], (70_000, 84_000)),
    "Fraudshield Data": ([("Transaction risk scoring API", .8), ("Model retraining", .2)], (18_000, 24_000)),
    "Pitchroom Events": ([("Fintech summit booth", .6), ("Lead scanning", .4)], (20_000, 30_000)),
    "Dealflow CRM": ([("Sales seats", .9), ("Forecasting module", .1)], (6_000, 6_500)),
    "Helix Cloud Health": ([("HIPAA-compliant compute", .6), ("Encrypted storage", .25), ("Audit logging", .15)], (40_000, 48_000)),
    "Clinitext AI": ([("Clinical note summarisation API", .7), ("Model fine-tuning", .3)], (9_000, 13_000)),
    "Brightpath Media": ([("Patient acquisition ads", .65), ("Landing page production", .35)], (11_000, 16_000)),
    "Caseline CRM": ([("Payer sales seats", .85), ("Contract module", .15)], (2_800, 3_000)),
}

PROFILES = {
    "kestrel": {
        "name": "Kestrel Metrics Inc.", "industry": "observability SaaS", "seed": 41, "peers": OBS_PEERS,
        "salary": {"engineering": 172_000, "ml_ai": 228_000, "sales": 149_000, "marketing": 138_000, "other": 118_000},
        "headcount": {"engineering": (64, 75), "ml_ai": (6, 11), "sales": (40, 47), "marketing": (9, 10), "other": (52, 55)},
        "infra_budget": 61_000, "tools": {"sales": 4_000, "marketing": 22_000, "ml_ai": 15_000},
        "budget_error_row": "sales",
        # (vendor, true category, cost centre printed, month, planted tag or None)
        "invoices": [("Stratus Hosting Co", "infra_cloud", "Infrastructure", m, None) for m in MONTHS[:5]] + [
            ("Quillfire Ads", "marketing", "Marketing", "2026-05", None),
            ("Quillfire Ads", "marketing", "Sales", "2026-08", "miscategorised"),
            ("Ledgerline CRM", "sales", "Sales", "2026-06", "totals_mismatch"),
            ("Vectorbay Compute", "ml_ai", "ML Platform", "2026-07", "scanned")],
        "near_miss": 3, "duplicate_of": 2, "zero_byte": "ledgerline_2026-09",
    },
    "ledgerly": {
        "name": "Ledgerly Payments Inc.", "industry": "fintech (spend management)", "seed": 53,
        "peers": {"competitors": [{"name": "Ramp", "domain": "ramp.com"}, {"name": "Brex", "domain": "brex.com"},
                                  {"name": "Mercury", "domain": "mercury.com"}, {"name": "Airwallex", "domain": "airwallex.com"},
                                  {"name": "Spendesk", "domain": "spendesk.com"}],
                  "substitutes": [{"name": "Pleo", "domain": "pleo.io"}, {"name": "Payhawk", "domain": "payhawk.com"}]},
        "salary": {"engineering": 190_000, "ml_ai": 245_000, "sales": 165_000, "marketing": 150_000, "other": 130_000},
        "headcount": {"engineering": (88, 102), "ml_ai": (9, 14), "sales": (95, 120), "marketing": (22, 26), "other": (70, 80)},
        "infra_budget": 78_000, "tools": {"sales": 6_500, "marketing": 28_000, "ml_ai": 22_000},
        "budget_error_row": "engineering",
        "invoices": [("Corvane Cloud", "infra_cloud", "Infrastructure", m, None) for m in MONTHS[:6]] + [
            ("Fraudshield Data", "ml_ai", "ML Platform", "2026-05", None),
            ("Fraudshield Data", "ml_ai", "Engineering", "2026-08", "miscategorised"),
            ("Pitchroom Events", "marketing", "Marketing", "2026-06", "scanned"),
            ("Dealflow CRM", "sales", "Sales", "2026-07", None),
            ("Dealflow CRM", "sales", "Sales", "2026-09", "totals_mismatch")],
        "near_miss": 10, "duplicate_of": 4, "zero_byte": None,
    },
    "carebridge": {
        "name": "Carebridge Health Inc.", "industry": "digital health", "seed": 67,
        "peers": {"competitors": [{"name": "Spring Health", "domain": "springhealth.com"},
                                  {"name": "Headway", "domain": "headway.co"},
                                  {"name": "Sword Health", "domain": "swordhealth.com"},
                                  {"name": "Cityblock Health", "domain": "cityblock.com"},
                                  {"name": "Komodo Health", "domain": "komodohealth.com"}],
                  "substitutes": [{"name": "Zocdoc", "domain": "zocdoc.com"}, {"name": "Oscar Health", "domain": "hioscar.com"}]},
        "salary": {"engineering": 178_000, "ml_ai": 232_000, "sales": 152_000, "marketing": 140_000, "other": 96_000},
        "headcount": {"engineering": (45, 52), "ml_ai": (4, 7), "sales": (18, 22), "marketing": (8, 10), "other": (160, 190)},
        "infra_budget": 44_000, "tools": {"sales": 3_000, "marketing": 15_000, "ml_ai": 11_000},
        "budget_error_row": "other",
        "invoices": [("Helix Cloud Health", "infra_cloud", "Infrastructure", m, None) for m in MONTHS[:5]] + [
            ("Clinitext AI", "ml_ai", "Marketing", "2026-06", "miscategorised_single"),  # vendor appears once
            ("Brightpath Media", "marketing", "Marketing", "2026-05", None),
            ("Brightpath Media", "marketing", "Marketing", "2026-08", "totals_mismatch"),
            ("Caseline CRM", "sales", "Sales", "2026-07", "scanned")],
        "near_miss": 2, "duplicate_of": None, "zero_byte": "helix_2026-10",
    },
}


def lines_for(vendor: str) -> list[tuple[str, int]]:
    items, (lo, hi) = VENDORS[vendor]
    base = random.randint(lo, hi)
    parts = [(d, round(base * s)) for d, s in items[:-1]]
    return parts + [(items[-1][0], base - sum(p for _, p in parts))]


def make_budget(p: dict, root: Path, fictional: str) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append([f"{p['name']} — FY27 people cost budget (USD)", "", fictional])
    ws.append(["category", "metric", *MONTHS, "FY total"])
    truth: dict = {"people_cost": {}, "headcount_end": {}}
    planted = None
    for cat, sal in p["salary"].items():
        a, b = p["headcount"][cat]
        hc = [round(a + (b - a) * i / 11) for i in range(12)]
        cost = [round(h * sal / 12) for h in hc]
        ws.append([cat, "headcount", *hc, ""])
        total = stated = sum(cost)
        if cat == p["budget_error_row"]:  # PLANTED: FY total typed 2.1% high
            stated = round(total * 1.021)
            planted = {"id": "budget_total_mismatch", "type": "totals_mismatch", "file": "budget_fy27.xlsx",
                       "sheet": "People", "row": ws.max_row + 1, "category": cat, "stated_total": stated,
                       "correct_total": total, "diff_pct": round(100 * (stated - total) / total, 2)}
        ws.append([cat, "people_cost", *cost, stated])
        truth["people_cost"][cat] = total
        truth["headcount_end"][cat] = hc[-1]
    ws2 = wb.create_sheet("NonPeople")
    ws2.append([f"{p['name']} — FY27 non-people budget (USD)", "", fictional])
    ws2.append(["category", "line", *MONTHS, "FY total"])
    truth["non_people_budget"] = {}
    for cat, line, amt in [("infra_cloud", "cloud + hosting", p["infra_budget"])] + \
                          [(c, "tools + programs", a) for c, a in p["tools"].items()]:
        ws2.append([cat, line, *[amt] * 12, amt * 12])
        truth["non_people_budget"][cat] = amt * 12
    wb.save(root / "budget_fy27.xlsx")
    return {"truth": truth, "planted": planted}


def make_invoices(p: dict, inv_dir: Path) -> dict:
    if inv_dir.exists():
        shutil.rmtree(inv_dir)
    inv_dir.mkdir(parents=True)
    prefix = "".join(w[0] for w in p["name"].split()[:2]).upper()
    invoices, planted, docs = [], [], []
    for n, (vendor, cat, cc, month, tag) in enumerate(p["invoices"], start=1):
        inv = {"number": f"{prefix}-{month.replace('-', '')}-{n:02d}", "vendor": vendor, "true_category": cat,
               "cost_centre": cc, "month": month, "lines": lines_for(vendor)}
        line_sum = sum(a for _, a in inv["lines"])
        stated, kind = line_sum, "text"
        fname = f"{n:02d}_{vendor.split()[0].lower()}_{month}.pdf"
        rel = f"invoices/{fname}"
        if tag == "totals_mismatch":
            stated = round(line_sum * 1.03)
            planted.append({"id": "invoice_total_mismatch", "type": "totals_mismatch", "file": rel,
                            "stated_total": stated, "line_items_sum": line_sum,
                            "diff_pct": round(100 * (stated - line_sum) / line_sum, 2)})
        if n == p["near_miss"]:  # NEGATIVE CONTROL: 0.4% rounding difference, inside the 1% tolerance
            stated = round(line_sum * 1.004)
            planted.append({"id": "near_miss_within_tolerance", "type": "negative_control", "file": rel,
                            "stated_total": stated, "line_items_sum": line_sum,
                            "expected": "NOT flagged (0.4% < 1% tolerance)"})
        if tag in ("miscategorised", "miscategorised_single"):
            planted.append({"id": "invoice_miscategorised" + ("_single_vendor" if tag.endswith("single") else ""),
                            "type": "miscategorised", "file": rel, "stated_cost_centre": cc,
                            "correct_category": cat, "amount": line_sum,
                            "vendor_appears_once": tag.endswith("single")})
        if tag == "scanned":
            kind = "scanned"
            write_scanned_pdf(inv_dir / fname, inv, stated)
            planted.append({"id": "invoice_scanned", "type": "scanned_pdf", "file": rel,
                            "expected": "text parser finds no text; LLM route extracts total", "amount": line_sum})
        else:
            write_text_pdf(inv_dir / fname, inv, stated)
        docs.append((fname, inv, stated))
        invoices.append({"file": rel, "number": inv["number"], "vendor": vendor, "kind": kind, "true_category": cat,
                         "cost_centre": cc, "month": month, "line_items_sum": line_sum, "stated_total": stated})
    nxt = len(p["invoices"]) + 1
    if p["duplicate_of"]:  # PLANTED: an invoice sent again under a new file name (same number, same amount)
        orig_name, orig, orig_stated = docs[p["duplicate_of"] - 1]
        dup = f"{nxt:02d}_{orig['vendor'].split()[0].lower()}_{orig['month']}_resent.pdf"
        write_text_pdf(inv_dir / dup, orig, orig_stated)
        planted.append({"id": "invoice_duplicate", "type": "duplicate", "file": f"invoices/{dup}",
                        "duplicate_of": f"invoices/{orig_name}", "number": orig["number"], "amount": orig_stated})
        nxt += 1
    if p["zero_byte"]:  # PLANTED: zero-byte PDF
        empty = f"{nxt:02d}_{p['zero_byte']}.pdf"
        (inv_dir / empty).write_bytes(b"")
        planted.append({"id": "invoice_corrupt", "type": "corrupt_file", "file": f"invoices/{empty}",
                        "expected": "zero bytes; listed under Not analysed"})
    true_by_cat: dict[str, int] = {}
    for inv in invoices:  # a duplicate is not real spend, so it is not counted
        true_by_cat[inv["true_category"]] = true_by_cat.get(inv["true_category"], 0) + inv["line_items_sum"]
    return {"invoices": invoices, "planted": planted, "invoice_totals_by_true_category": true_by_cat}


def build(key: str) -> None:
    p = PROFILES[key]
    random.seed(p["seed"])
    root = Path("data") / key
    root.mkdir(parents=True, exist_ok=True)
    fictional = f"FICTIONAL SAMPLE DATA — {p['name']} does not exist."
    budget = make_budget(p, root, fictional)
    inv = make_invoices(p, root / "invoices")
    policy = yaml.safe_load(Path("data/acme/policy.yaml").read_text(encoding="utf-8"))
    policy["_note"] = fictional
    (root / "policy.yaml").write_text(yaml.safe_dump(policy, sort_keys=False), encoding="utf-8")
    peers = {"_note": f"Real companies; data from Crustdata. {p['name']} itself is FICTIONAL.", **p["peers"]}
    (root / "competitors.yaml").write_text(yaml.safe_dump(peers, sort_keys=False), encoding="utf-8")
    with open(root / "salary_table.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["category", "loaded_annual_cost_usd", "source"])
        w.writerows([c, v, "ASSUMPTION (illustrative, not survey data)"] for c, v in p["salary"].items())
    (root / "README.md").write_text(f"# {fictional}\n\nHeld-out eval company ({p['industry']}) generated by "
                                    "`scripts/make_heldout_data.py`. The agent was not tuned on it.\n", encoding="utf-8")
    truth = {
        "_note": "HELD-OUT eval ground truth. NEVER shown to the agent. Generated, not hand-typed.",
        "company": p["name"], "industry": p["industry"],
        "people_cost_fy27": budget["truth"]["people_cost"],
        "headcount_end_fy27": budget["truth"]["headcount_end"],
        "non_people_budget_fy27": budget["truth"]["non_people_budget"],
        "invoice_totals_by_true_category": inv["invoice_totals_by_true_category"],
        "invoices": inv["invoices"],
        "planted_errors": [budget["planted"], *inv["planted"]],
        "fixtures": [],
    }
    Path(f"evals/ground_truth_{key}.json").write_text(json.dumps(truth, indent=2), encoding="utf-8")
    n_files = len(list((root / "invoices").iterdir()))
    print(f"{key}: {p['name']} ({p['industry']}), {n_files} invoice files, "
          f"{len(truth['planted_errors'])} planted items: {[e['id'] for e in truth['planted_errors']]}")


if __name__ == "__main__":
    for k in sys.argv[1:] or PROFILES:
        build(k)
