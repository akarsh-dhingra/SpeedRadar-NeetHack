"""Generate the FICTIONAL Acme Observability Inc. finance folder + eval ground truth.

Deterministic (fixed seed). Writes:
  data/acme/budget_fy27.xlsx      people + non-people budget by category x month
  data/acme/invoices/*.pdf        13 invoices, incl. planted problems
  data/acme/policy.yaml           people-cost policy vs competitor median
  data/acme/salary_table.csv      loaded cost per role category (ASSUMPTION)
  evals/ground_truth.json         planted errors + correct totals (never shown to the agent)

Vendors are invented names. Run: uv run python scripts/make_acme_data.py
"""

from __future__ import annotations

import csv
import io
import json
import random
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font

from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

random.seed(27)
ROOT = Path("data/acme")
INV = ROOT / "invoices"
FICTIONAL = "FICTIONAL SAMPLE DATA — Acme Observability Inc. does not exist."
MONTHS = ["2026-04", "2026-05", "2026-06", "2026-07", "2026-08", "2026-09",
          "2026-10", "2026-11", "2026-12", "2027-01", "2027-02", "2027-03"]

# Loaded annual cost per head, USD. ASSUMPTION for the demo, not survey data.
SALARY = {"engineering": 185_000, "ml_ai": 240_000, "sales": 155_000, "marketing": 145_000}
# Headcount plan: start of FY → end of FY (linear hiring).
HEADCOUNT = {"engineering": (118, 130), "ml_ai": (14, 22), "sales": (58, 64), "marketing": (21, 24)}
INFRA_MONTHLY_BUDGET = 92_000
TOOLS_MONTHLY_BUDGET = {"sales": 6_000, "marketing": 38_000, "ml_ai": 30_000}


def headcount_for(cat: str, i: int) -> int:
    a, b = HEADCOUNT[cat]
    return round(a + (b - a) * i / 11)


def make_budget() -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["Acme Observability Inc. — FY27 people cost budget (USD)", "", FICTIONAL])
    ws.append(["category", "metric", *MONTHS, "FY total"])
    for c in ws[2]:
        c.font = Font(bold=True)
    truth: dict = {"people_cost": {}, "headcount_end": {}}
    planted = None
    for cat in SALARY:
        hc = [headcount_for(cat, i) for i in range(12)]
        cost = [round(h * SALARY[cat] / 12) for h in hc]
        ws.append([cat, "headcount", *hc, ""])
        total = sum(cost)
        stated = total
        if cat == "marketing":  # PLANTED: stated total does not equal the sum of months
            stated = total + 48_000
            planted = {"id": "budget_total_mismatch", "type": "totals_mismatch",
                       "file": "budget_fy27.xlsx", "sheet": "People", "row": ws.max_row + 1,
                       "category": "marketing", "stated_total": stated, "correct_total": total,
                       "diff_pct": round(100 * (stated - total) / total, 2)}
        ws.append([cat, "people_cost", *cost, stated])
        truth["people_cost"][cat] = total
        truth["headcount_end"][cat] = hc[-1]

    ws2 = wb.create_sheet("NonPeople")
    ws2.append(["Acme Observability Inc. — FY27 non-people budget (USD)", "", FICTIONAL])
    ws2.append(["category", "line", *MONTHS, "FY total"])
    for c in ws2[2]:
        c.font = Font(bold=True)
    rows = [("infra_cloud", "cloud compute + storage", INFRA_MONTHLY_BUDGET)]
    rows += [(cat, "tools + programs", amt) for cat, amt in TOOLS_MONTHLY_BUDGET.items()]
    truth["non_people_budget"] = {}
    for cat, line, amt in rows:
        vals = [amt] * 12
        ws2.append([cat, line, *vals, sum(vals)])
        truth["non_people_budget"][cat] = sum(vals)
    ROOT.mkdir(parents=True, exist_ok=True)
    wb.save(ROOT / "budget_fy27.xlsx")
    return {"truth": truth, "planted": planted}


def invoice_lines(vendor: str, cat: str, month: str) -> list[tuple[str, int]]:
    if vendor == "Nimbus Cloud Ltd":
        base = INFRA_MONTHLY_BUDGET + random.randint(-9_000, 9_000)
        parts = [("Compute instances", round(base * 0.58)), ("Block + object storage", round(base * 0.22)),
                 ("Network egress", round(base * 0.12))]
        parts.append(("Managed database", base - sum(p for _, p in parts)))
        return parts
    if vendor == "Tensorforge GPU":
        base = 26_000 + random.randint(0, 8_000)
        return [("A-class GPU hours (training)", round(base * 0.8)), ("Inference endpoints", base - round(base * 0.8))]
    if vendor == "Brightloop Media":
        base = 30_000 + random.randint(0, 10_000)
        return [("Paid search management", round(base * 0.6)), ("Conference sponsorship", base - round(base * 0.6))]
    base = 5_400
    return [("CRM seats (60)", 4_200), ("Sales engagement add-on", base - 4_200)]


def draw_invoice(c: canvas.Canvas, inv: dict, stated_total: int) -> None:
    w, h = A4
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, h - 60, inv["vendor"])
    c.setFont("Helvetica", 9)
    c.drawString(50, h - 75, FICTIONAL)
    c.setFont("Helvetica", 11)
    y = h - 110
    for label, val in [("Invoice no.", inv["number"]), ("Invoice date", f"{inv['month']}-28"),
                       ("Bill to", "Acme Observability Inc."), ("Cost centre", inv["cost_centre"]),
                       ("Currency", "USD")]:
        c.drawString(50, y, f"{label}: {val}")
        y -= 16
    y -= 14
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "Description")
    c.drawRightString(540, y, "Amount")
    c.setFont("Helvetica", 11)
    for desc, amt in inv["lines"]:
        y -= 18
        c.drawString(50, y, desc)
        c.drawRightString(540, y, f"{amt:,.2f}")
    y -= 26
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, y, "TOTAL DUE")
    c.drawRightString(540, y, f"{stated_total:,.2f}")


def write_text_pdf(path: Path, inv: dict, stated_total: int) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    draw_invoice(c, inv, stated_total)
    c.showPage()
    c.save()


def write_scanned_pdf(path: Path, inv: dict, stated_total: int) -> None:
    """Render the invoice to an image and embed only the image: no text layer."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    draw_invoice(c, inv, stated_total)
    c.showPage()
    c.save()
    import pypdfium2 as pdfium

    page = pdfium.PdfDocument(buf.getvalue())[0]
    img = page.render(scale=1.6).to_pil().convert("L").rotate(0.7, fillcolor=255, expand=False)
    img = img.point(lambda p: 255 if p > 200 else max(0, p - 20))  # photocopy look
    out = canvas.Canvas(str(path), pagesize=A4)
    out.drawImage(ImageReader(img), 0, 0, width=A4[0], height=A4[1])
    out.showPage()
    out.save()


def make_invoices() -> dict:
    INV.mkdir(parents=True, exist_ok=True)
    plan = []
    for i, m in enumerate(MONTHS[:6]):
        plan.append(("Nimbus Cloud Ltd", "infra_cloud", "Infrastructure", m))
    plan += [("Tensorforge GPU", "ml_ai", "ML Platform", "2026-05"),
             ("Tensorforge GPU", "ml_ai", "Marketing", "2026-08"),  # PLANTED: mis-categorised
             ("Brightloop Media", "marketing", "Marketing", "2026-06"),
             ("Brightloop Media", "marketing", "Marketing", "2026-09"),
             ("Pipewise CRM", "sales", "Sales", "2026-07"),
             ("Nimbus Cloud Ltd", "infra_cloud", "Infrastructure", "2026-10")]  # scanned
    invoices, planted = [], []
    for n, (vendor, cat, cc, month) in enumerate(plan, start=1):
        number = f"INV-{month.replace('-', '')}-{n:03d}"
        inv = {"number": number, "vendor": vendor, "true_category": cat, "cost_centre": cc,
               "month": month, "lines": invoice_lines(vendor, cat, month)}
        line_sum = sum(a for _, a in inv["lines"])
        stated = line_sum
        fname = f"{n:02d}_{vendor.split()[0].lower()}_{month}.pdf"
        kind = "text"
        if n == 4:  # PLANTED: line items do not sum to the stated total
            stated = line_sum + 7_500
            planted.append({"id": "invoice_total_mismatch", "type": "totals_mismatch", "file": f"invoices/{fname}",
                            "stated_total": stated, "line_items_sum": line_sum,
                            "diff_pct": round(100 * (stated - line_sum) / line_sum, 2)})
        if cc == "Marketing" and cat == "ml_ai":
            planted.append({"id": "invoice_miscategorised", "type": "miscategorised", "file": f"invoices/{fname}",
                            "stated_cost_centre": cc, "correct_category": cat, "amount": line_sum})
        if n == len(plan):
            kind = "scanned"
            write_scanned_pdf(INV / fname, inv, stated)
            planted.append({"id": "invoice_scanned", "type": "scanned_pdf", "file": f"invoices/{fname}",
                            "expected": "text parser finds no text; LLM/vision route extracts total",
                            "amount": line_sum})
        else:
            write_text_pdf(INV / fname, inv, stated)
        invoices.append({"file": f"invoices/{fname}", "number": number, "vendor": vendor, "kind": kind,
                         "true_category": cat, "cost_centre": cc, "month": month,
                         "line_items_sum": line_sum, "stated_total": stated})
    # PLANTED: corrupt file — a real PDF header followed by truncated garbage
    corrupt = "13_brightloop_2026-10.pdf"
    (INV / corrupt).write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
                                + bytes(random.getrandbits(8) for _ in range(600)))
    planted.append({"id": "invoice_corrupt", "type": "corrupt_file", "file": f"invoices/{corrupt}",
                    "expected": "both parsers fail; listed under Not analysed"})
    true_by_cat: dict[str, int] = {}
    for inv in invoices:
        true_by_cat[inv["true_category"]] = true_by_cat.get(inv["true_category"], 0) + inv["line_items_sum"]
    return {"invoices": invoices, "planted": planted, "invoice_totals_by_true_category": true_by_cat}


def make_policy() -> None:
    policy = {
        "_note": FICTIONAL,
        "basis": "people cost (headcount x loaded cost), FY27, compared with the competitor median per head of company",
        "rules": [
            {"id": "ml_ai_parity", "category": "ml_ai", "metric": "share_of_headcount",
             "rule": "within_pct_of_median", "pct": 20,
             "text": "ML/AI share of headcount within 20% of competitor median"},
            {"id": "eng_floor", "category": "engineering", "metric": "share_of_headcount",
             "rule": "at_least_pct_of_median", "pct": 90,
             "text": "Engineering share of headcount at least 90% of competitor median"},
            {"id": "sales_band", "category": "sales", "metric": "share_of_headcount",
             "rule": "within_pct_of_median", "pct": 25,
             "text": "Sales share of headcount within 25% of competitor median"},
            {"id": "marketing_floor", "category": "marketing", "metric": "share_of_headcount",
             "rule": "at_least_pct_of_median", "pct": 80,
             "text": "Marketing share of headcount at least 80% of competitor median"},
            {"id": "infra_budget", "category": "infra_cloud", "metric": "invoice_vs_budget",
             "rule": "max_over_budget_pct", "pct": 5,
             "text": "Infra/cloud invoices no more than 5% over monthly budget (internal check only)"},
        ],
    }
    (ROOT / "policy.yaml").write_text(yaml.safe_dump(policy, sort_keys=False))


def make_salary_table() -> None:
    with open(ROOT / "salary_table.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["category", "loaded_annual_cost_usd", "source"])
        for cat, v in SALARY.items():
            w.writerow([cat, v, "ASSUMPTION (illustrative, not survey data)"])


def main() -> None:
    budget = make_budget()
    inv = make_invoices()
    make_policy()
    make_salary_table()
    (ROOT / "README.md").write_text(
        f"# {FICTIONAL}\n\nGenerated by `scripts/make_acme_data.py` (seeded, reproducible). "
        "Vendor names are invented. Competitors in `competitors.yaml` are real companies; their data "
        "comes live from Crustdata.\n")
    truth = {
        "_note": "Eval ground truth. NEVER shown to the agent. Generated, not hand-typed.",
        "people_cost_fy27": budget["truth"]["people_cost"],
        "headcount_end_fy27": budget["truth"]["headcount_end"],
        "non_people_budget_fy27": budget["truth"]["non_people_budget"],
        "invoice_totals_by_true_category": inv["invoice_totals_by_true_category"],
        "invoices": inv["invoices"],
        "planted_errors": [budget["planted"], *inv["planted"]],
        "fixtures": [],  # fault-injected competitor fixtures are added by evals/make_fixtures.py
    }
    Path("evals").mkdir(exist_ok=True)
    Path("evals/ground_truth.json").write_text(json.dumps(truth, indent=2))
    print(f"wrote {len(inv['invoices']) + 1} invoices, {len(truth['planted_errors'])} planted errors")


if __name__ == "__main__":
    main()
